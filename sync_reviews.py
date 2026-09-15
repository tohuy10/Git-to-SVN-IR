#!/usr/bin/env python3
import os
import re
import sys
import subprocess
from datetime import datetime, timedelta, timezone
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ==============================================================================
# CONFIGURATION
# ==============================================================================
EXCEL_FILE_NAME = "Git_Review_Log.xlsx"
SHEET_NAME = "Review Logs"
TZ_GMT7 = timezone(timedelta(hours=7))
DATE_DISPLAY_FORMAT = "%d %b, %Y %I:%M %p"  # e.g., 11 Sep, 2026 12:08 AM
commit_cache = {}


def normalize_repo_slug(raw: str) -> str:
    cleaned = raw.strip()
    match = re.search(r"github\.com[/:]([\w-]+/[\w.-]+?)(?:\.git)?$", cleaned)
    return match.group(1) if match else cleaned


def run_cmd(cmd, cwd=None):
    print(f"[RUNNING] {' '.join(cmd)}", flush=True)
    result = subprocess.run(cmd, cwd=cwd, text=True, capture_output=True)
    if result.returncode != 0:
        print(f"[CMD ERROR] {' '.join(cmd)}\nSTDERR: {result.stderr.strip()}", file=sys.stderr, flush=True)
        raise subprocess.CalledProcessError(result.returncode, cmd, result.stdout, result.stderr)
    return result.stdout.strip()


def sync_svn(work_dir: str, svn_url: str, user: str, password: str):
    auth_args = [
        "--non-interactive",
        "--no-auth-cache",
        "--trust-server-cert",
        "--trust-server-cert-failures=unknown-ca,cn-mismatch,expired,not-yet-valid,other"
    ]
    if user:
        auth_args.extend(["--username", user, "--password", password])

    needs_fresh_checkout = True

    if os.path.exists(os.path.join(work_dir, ".svn")):
        try:
            current_wc_url = run_cmd(["svn", "info", "--show-item", "url", work_dir] + auth_args)
            if current_wc_url.rstrip("/") == svn_url.rstrip("/"):
                needs_fresh_checkout = False
            else:
                print(f"[INFO] SVN URL changed from '{current_wc_url}' to '{svn_url}'. Performing clean checkout.", flush=True)
        except Exception:
            needs_fresh_checkout = True

    if needs_fresh_checkout:
        if os.path.exists(work_dir):
            import shutil
            shutil.rmtree(work_dir)
        cmd = ["svn", "checkout", svn_url, work_dir] + auth_args
        print(f"Executing: svn checkout {svn_url} {work_dir}", flush=True)
        run_cmd(cmd)
    else:
        cmd = ["svn", "update", work_dir] + auth_args
        print("Executing: svn update", flush=True)
        run_cmd(cmd)


def commit_to_svn(work_dir: str, user: str, password: str):
    auth_args = [
        "--non-interactive",
        "--no-auth-cache",
        "--trust-server-cert",
        "--trust-server-cert-failures=unknown-ca,cn-mismatch,expired,not-yet-valid,other"
    ]
    if user:
        auth_args.extend(["--username", user, "--password", password])

    status = run_cmd(["svn", "status", EXCEL_FILE_NAME], cwd=work_dir)
    if not status:
        print("No changes detected in Excel log. Skipping SVN commit.", flush=True)
        return

    if status.startswith("?"):
        run_cmd(["svn", "add", "--force", EXCEL_FILE_NAME], cwd=work_dir)

    commit_cmd = ["svn", "commit", "-m", "[Auto-Sync] Update Git review logs via Jenkins", EXCEL_FILE_NAME] + auth_args
    output = run_cmd(commit_cmd, cwd=work_dir)
    print(output, flush=True)


def get_commit_details(repo: str, sha: str, headers: dict):
    if not sha:
        return {}
    if sha in commit_cache:
        return commit_cache[sha]

    url = f"https://api.github.com/repos/{repo}/commits/{sha}"
    resp = requests.get(url, headers=headers, timeout=15)
    if resp.status_code == 200:
        data = resp.json()
        commit_cache[sha] = data
        return data
    return {}


def parse_github_datetime(iso_str: str):
    if not iso_str:
        return datetime.min.replace(tzinfo=timezone.utc), ""
    utc_dt = datetime.fromisoformat(iso_str.replace("Z", "+00:00"))
    gmt7_dt = utc_dt.astimezone(TZ_GMT7)
    return utc_dt, gmt7_dt.strftime(DATE_DISPLAY_FORMAT)


def fetch_all_pages(base_url: str, headers: dict) -> list:
    items = []
    current_url = f"{base_url}{'&' if '?' in base_url else '?'}per_page=100"

    while current_url:
        resp = requests.get(current_url, headers=headers, timeout=15)
        if resp.status_code != 200:
            break
        data = resp.json()
        if not isinstance(data, list):
            break
        items.extend(data)
        current_url = resp.links.get("next", {}).get("url")

    return items


def get_repo_watermarks(file_path: str) -> dict:
    watermarks = {}
    if not os.path.exists(file_path):
        return watermarks

    supported_formats = [
        DATE_DISPLAY_FORMAT,
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
    ]

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 4:
                continue
            repo = str(row[0] or "").strip()
            created_at_val = row[3]
            if not repo or not created_at_val:
                continue

            dt_gmt7 = None
            if isinstance(created_at_val, datetime):
                if created_at_val.tzinfo is None:
                    dt_gmt7 = created_at_val.replace(tzinfo=TZ_GMT7)
                else:
                    dt_gmt7 = created_at_val.astimezone(TZ_GMT7)
            elif isinstance(created_at_val, str):
                cleaned_str = created_at_val.strip()
                for fmt in supported_formats:
                    try:
                        parsed = datetime.strptime(cleaned_str, fmt)
                        dt_gmt7 = parsed.replace(tzinfo=TZ_GMT7)
                        break
                    except ValueError:
                        continue

            if dt_gmt7:
                utc_dt = dt_gmt7.astimezone(timezone.utc)
                if repo not in watermarks or utc_dt > watermarks[repo]:
                    watermarks[repo] = utc_dt

        wb.close()
    except Exception as e:
        print(f"[WARN] Error reading watermarks from {file_path}: {e}", flush=True)

    return watermarks


def compute_cutoff_utc(repo: str, scan_timeframe: str, watermark_utc: datetime | None) -> datetime | None:
    now_utc = datetime.now(timezone.utc)
    mode = (scan_timeframe or "").strip().lower()

    # 1. Full history bypass
    if "all" in mode or "every" in mode:
        print(f"[{repo}] Mode: 'All History'. Scanning all Pull Requests...", flush=True)
        return None

    # 2. Incremental Excel watermark check
    if "incremental" in mode or "excel" in mode:
        if watermark_utc:
            cutoff = watermark_utc - timedelta(minutes=10)
            cutoff_disp = cutoff.astimezone(TZ_GMT7).strftime(DATE_DISPLAY_FORMAT)
            print(f"[{repo}] Mode: 'Incremental'. Watermark found in Excel: fetching PRs active after {cutoff_disp} GMT+7", flush=True)
            return cutoff
        else:
            print(f"[{repo}] Mode: 'Incremental' selected, but no Excel record exists. Scanning all history...", flush=True)
            return None

    # 3. Dynamic Relative Timeframes (1-3 weeks, 1-6 months, 1-2 years, etc.)
    match = re.search(r"(\d+)\s*(day|week|month|year)", mode)
    if match:
        count = int(match.group(1))
        unit = match.group(2)
        if unit == "day":
            days = count
        elif unit == "week":
            days = count * 7
        elif unit == "month":
            days = count * 30
        elif unit == "year":
            days = count * 365
        cutoff = now_utc - timedelta(days=days)
    else:
        # Fallback default: 1 week (7 days)
        cutoff = now_utc - timedelta(days=7)

    cutoff_disp = cutoff.astimezone(TZ_GMT7).strftime(DATE_DISPLAY_FORMAT)
    print(f"[{repo}] Mode: '{scan_timeframe}'. Scanning PRs active after {cutoff_disp} GMT+7", flush=True)
    return cutoff


def process_repository(repo: str, token: str, watermark_utc: datetime | None, scan_timeframe: str):
    headers = {
        "Authorization": f"Bearer {token}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    }
    all_repo_records = []

    cutoff_utc = compute_cutoff_utc(repo, scan_timeframe, watermark_utc)

    page = 1
    stop_pagination = False
    pr_per_page = 100

    while True:
        print(f"[{repo}] Fetching PR page {page} ({pr_per_page} PRs per page)...", flush=True)
        prs_url = f"https://api.github.com/repos/{repo}/pulls?state=all&sort=updated&direction=desc&per_page={pr_per_page}&page={page}"
        resp = requests.get(prs_url, headers=headers, timeout=20)
        if resp.status_code != 200:
            print(f"Failed to fetch PRs for {repo} (page {page}): {resp.status_code} - {resp.text}", flush=True)
            break

        prs = resp.json()
        if not prs or not isinstance(prs, list):
            break

        print(f"[{repo}] Page {page}: Processing {len(prs)} PRs...", flush=True)

        for pr in prs:
            pr_updated_at_str = pr.get("updated_at")
            if pr_updated_at_str and cutoff_utc:
                pr_up_utc = datetime.fromisoformat(pr_updated_at_str.replace("Z", "+00:00"))
                if pr_up_utc < cutoff_utc:
                    stop_pagination = True
                    break

            pr_number = pr["number"]
            default_dev = pr.get("user", {}).get("login", "")

            # 1. Main PR Reviews (/reviews)
            reviews = fetch_all_pages(f"https://api.github.com/repos/{repo}/pulls/{pr_number}/reviews", headers)
            for rev in reviews:
                body = (rev.get("body") or "").strip()
                if not body:
                    continue

                sub_utc, display_gmt7 = parse_github_datetime(rev.get("submitted_at"))
                if cutoff_utc and sub_utc < cutoff_utc:
                    continue

                commit_sha = rev.get("commit_id") or pr.get("head", {}).get("sha", "")
                commit_info = get_commit_details(repo, commit_sha, headers)

                dev_name = (
                    commit_info.get("commit", {}).get("author", {}).get("name")
                    or default_dev
                )
                commit_msg = commit_info.get("commit", {}).get("message", "")

                all_repo_records.append({
                    "raw_dt": sub_utc,
                    "repo_name": repo,
                    "commit_id": commit_sha,
                    "dev_name": dev_name,
                    "created_at": display_gmt7,
                    "commit_content": commit_msg,
                    "reviewer_name": rev.get("user", {}).get("login", ""),
                    "reviewer_comment": body,
                    "review_type": "PR Review",
                    "file_path_line": "—",
                    "review_state": rev.get("state", ""),
                })

            # 2. Line-Specific Comments (/comments)
            comments = fetch_all_pages(f"https://api.github.com/repos/{repo}/pulls/{pr_number}/comments", headers)
            for c in comments:
                c_utc, display_gmt7 = parse_github_datetime(c.get("created_at"))
                if cutoff_utc and c_utc < cutoff_utc:
                    continue

                commit_sha = c.get("commit_id") or pr.get("head", {}).get("sha", "")
                commit_info = get_commit_details(repo, commit_sha, headers)

                dev_name = (
                    commit_info.get("commit", {}).get("author", {}).get("name")
                    or default_dev
                )
                commit_msg = commit_info.get("commit", {}).get("message", "")

                path = c.get("path", "")

                # Fall back to original_line if line is null (e.g. outdated/resolved comments)
                # can choose to not fall back to original line number if outdated comment line number is not useful
                line = c.get("line") or c.get("original_line")
                start_line = c.get("start_line") or c.get("original_start_line")

                if start_line and line and start_line != line:
                    loc = f"{path}:L{start_line}-L{line}"
                elif line:
                    loc = f"{path}:L{line}"
                else:
                    loc = path

                all_repo_records.append({
                    "raw_dt": c_utc,
                    "repo_name": repo,
                    "commit_id": commit_sha,
                    "dev_name": dev_name,
                    "created_at": display_gmt7,
                    "commit_content": commit_msg,
                    "reviewer_name": c.get("user", {}).get("login", ""),
                    "reviewer_comment": c.get("body", ""),
                    "review_type": "Line Specific",
                    "file_path_line": loc,
                    "review_state": "—",
                })

        print(f"[{repo}] Page {page} done. Total reviews collected so far: {len(all_repo_records)}", flush=True)

        if stop_pagination or len(prs) < pr_per_page or "next" not in resp.links:
            break
        page += 1

    return all_repo_records


def update_excel(file_path: str, records: list):
    headers = [
        "Repository", "Commit ID", "Dev name", "Created at", "Commit content",
        "Reviewer name", "Reviewer comment", "Review Type", "File Path / Line", "Review State"
    ]
    existing_keys = set()

    if os.path.exists(file_path):
        wb = openpyxl.load_workbook(file_path)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and len(row) >= 7:
                created_raw = row[3]
                if isinstance(created_raw, datetime):
                    created_str = created_raw.strftime(DATE_DISPLAY_FORMAT)
                else:
                    created_str = str(created_raw or "").strip()

                comment_text = str(row[6] or "").strip()
                key = f"{row[0]}|{row[1]}|{row[5]}|{created_str}|{comment_text}"
                existing_keys.add(key)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(headers)

        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_align = Alignment(horizontal="center", vertical="center")

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
        ws.row_dimensions[1].height = 25

    data_align = Alignment(vertical="top", wrap_text=True)
    added_count = 0

    for rec in records:
        clean_comment = rec["reviewer_comment"].strip()
        key = f"{rec['repo_name']}|{rec['commit_id']}|{rec['reviewer_name']}|{rec['created_at']}|{clean_comment}"
        if key in existing_keys:
            continue
        existing_keys.add(key)

        row_data = [
            rec["repo_name"],
            rec["commit_id"],
            rec["dev_name"],
            rec["created_at"],
            rec["commit_content"],
            rec["reviewer_name"],
            rec["reviewer_comment"].replace("\r\n", "\n"),
            rec["review_type"],
            rec["file_path_line"],
            rec["review_state"],
        ]

        ws.append(row_data)
        current_row = ws.max_row
        for col_idx in range(1, len(row_data) + 1):
            ws.cell(row=current_row, column=col_idx).alignment = data_align
        added_count += 1

    col_widths = {
        "A": 28, "B": 22, "C": 16, "D": 24, "E": 30,
        "F": 16, "G": 45, "H": 15, "I": 22, "J": 14
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(file_path)
    print(f"Excel sync complete: {added_count} new records added.", flush=True)


def main():
    github_token = (os.getenv("GITHUB_TOKEN") or "").strip()
    repo_list_env = (os.getenv("REPO_LIST") or "").strip()
    svn_url = (os.getenv("SVN_URL") or "").strip()
    svn_user = (os.getenv("SVN_USER") or "").strip()
    svn_pass = (os.getenv("SVN_PASS") or "").strip()
    scan_timeframe = (os.getenv("SCAN_TIMEFRAME") or "1 week").strip()

    if not github_token:
        print("[ERROR] GITHUB_TOKEN environment variable not set.", file=sys.stderr)
        sys.exit(1)
    if not repo_list_env:
        print("[ERROR] REPO_LIST environment variable not set.", file=sys.stderr)
        sys.exit(1)
    if not svn_url:
        print("[ERROR] SVN_URL environment variable not set.", file=sys.stderr)
        sys.exit(1)

    work_dir = "svn_workdir"

    print("==> Step 1: Syncing SVN Working Directory...", flush=True)
    sync_svn(work_dir, svn_url, svn_user, svn_pass)

    excel_path = os.path.join(work_dir, EXCEL_FILE_NAME)
    watermarks = {}

    print(f"==> Step 2: Evaluating scan timeframe: '{scan_timeframe}'...", flush=True)
    if "incremental" in scan_timeframe.lower() or "excel" in scan_timeframe.lower():
        watermarks = get_repo_watermarks(excel_path)
        for r, wm in watermarks.items():
            wm_display = wm.astimezone(TZ_GMT7).strftime(DATE_DISPLAY_FORMAT)
            print(f"    - {r}: Latest review timestamp = {wm_display} (GMT+7)", flush=True)

    print("==> Step 3: Fetching GitHub Review Logs...", flush=True)
    all_records = []

    for repo_raw in repo_list_env.split(","):
        repo = normalize_repo_slug(repo_raw)
        if not repo:
            continue
        print(f"Processing repository: {repo}", flush=True)
        repo_watermark = watermarks.get(repo)
        records = process_repository(repo, github_token, repo_watermark, scan_timeframe)
        all_records.extend(records)

    print(f"==> Step 4: Sorting {len(all_records)} total review records chronologically (GMT+7)...", flush=True)
    all_records.sort(key=lambda item: item["raw_dt"])

    print(f"==> Step 5: Updating Excel file at {excel_path}...", flush=True)
    update_excel(excel_path, all_records)

    print("==> Step 6: Committing updated Excel to SVN...", flush=True)
    commit_to_svn(work_dir, svn_user, svn_pass)

    print("[SUCCESS] Process completed successfully.", flush=True)


if __name__ == "__main__":
    main()